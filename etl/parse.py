"""P1: 第1・5・6表のパーサー（Excel: 2022-2024 / PDF: 2011-2021）。

値は公表値のまま（Excelの按分小数を含む）。欠損は None。
"""
import re
import unicodedata
from pathlib import Path

import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"

REGIONS = ["下越", "中越", "魚沼", "上越", "佐渡"]
CATEGORIES = ["自然", "歴史・文化", "温泉・健康", "スポーツ・レクリエーション",
              "都市型観光", "その他", "行祭事・イベント"]

EXCEL_FILES = {2024: "483582.xlsx", 2023: "483583.xlsx", 2022: "407412.xlsx"}
PDF_FILES = {2021: "343294.pdf", 2020: "293902.pdf", 2019: "240907.pdf", 2018: "195170.pdf",
             2017: "50508.pdf", 2016: "50376.pdf", 2015: "50560.pdf", 2014: "50370.pdf",
             2013: "50339.pdf", 2012: "50581.pdf", 2011: "50453.pdf"}
YEARS = sorted(list(EXCEL_FILES) + list(PDF_FILES))

_INT = re.compile(r"^[0-9][0-9,]*(\.[0-9]+)?$")


def nfkc(s):
    return unicodedata.normalize("NFKC", s or "")


def compact(s):
    return re.sub(r"[\s　]", "", nfkc(s))


def era_to_year(text):
    t = nfkc(text)
    if "令和元" in t or "平成31" in t and "令和" in t:
        return 2019
    m = re.search(r"(令和|平成)\s*(\d+)\s*年", t)
    if not m:
        return None
    n = int(m.group(2))
    return 2018 + n if m.group(1) == "令和" else 1988 + n


def to_num(v):
    """数値化。欠損（－・空）は None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = compact(str(v)).replace(",", "")
    if s in ("", "－", "-", "―", "—"):
        return None
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def norm_muni(name):
    """市町村名の正規化。（※）等を除去し、佐渡は『佐渡市』に統一。集計行は ('region', 地域名)/('total',) を返す。"""
    n = compact(name).replace("(※)", "").replace("※", "")
    if n.startswith("合計"):
        return ("total", None)
    m = re.match(r"(下越|中越|魚沼|上越)(地域|地区)計", n)
    if m:
        return ("region", m.group(1))
    if n.startswith("佐渡"):
        if n.endswith("計"):  # 『佐渡地域計』(2011-2020前後)
            return ("region", "佐渡")
        return ("muni", "佐渡市")
    return ("muni", n)


def _group_regions(rows):
    """rows: [(kind, key, payload)] → 地域計行に出会ったら直前の市町村群に地域を付ける。"""
    out, buf = [], []
    for kind, key, payload in rows:
        if kind == "muni":
            buf.append((key, payload))
        elif kind == "region":
            if key == "佐渡" and not buf:  # 佐渡地域 計 が単独行（2011〜）
                out.append(("佐渡", "佐渡市", payload))
                continue
            for m, p in buf:
                out.append((key, m, p))
            buf = []
            out.append(("_region_total", key, payload))
        else:
            out.append(("_total", "合計", payload))
    # 佐渡市（1行）は地域計行を持たない
    for m, p in buf:
        out.append(("佐渡" if m == "佐渡市" else "?", m, p))
    return out


# ---------------- Excel ----------------
def _xl_sheet(year, prefix):
    wb = openpyxl.load_workbook(RAW / str(year) / EXCEL_FILES[year], data_only=True)
    for ws in wb:
        if ws.title.startswith(prefix):
            return ws
    raise KeyError(prefix)


def excel_table1(year):
    ws = _xl_sheet(year, "第１表")
    out = {}
    for r in ws.iter_rows(min_row=5, max_col=3, values_only=True):
        y = era_to_year(str(r[0] or ""))
        if y and isinstance(r[1], (int, float)):
            out[y] = (float(r[1]), float(r[2]) if r[2] is not None else None)
    return out


def excel_table5(year):
    ws = _xl_sheet(year, "表5")
    rows = []
    for r in ws.iter_rows(min_row=5, max_col=14, values_only=True):
        if not r[0]:
            continue
        kind, key = norm_muni(str(r[0]))
        if not any(isinstance(v, (int, float)) for v in r[1:14]):
            continue
        months = [to_num(v) for v in r[1:13]]
        rows.append((kind, key, {"months": months, "total": to_num(r[13])}))
    return _group_regions(rows)


def excel_table6(year):
    ws = _xl_sheet(year, "表6")
    rows = []
    for r in ws.iter_rows(min_row=5, max_col=9, values_only=True):
        if not r[0] or not any(isinstance(v, (int, float)) for v in r[1:9]):
            continue
        kind, key = norm_muni(str(r[0]))
        rows.append((kind, key, {"cats": [to_num(v) for v in r[1:8]], "total": to_num(r[8])}))
    return _group_regions(rows)


# ---------------- PDF ----------------
_pdf_cache = {}


def pdf_texts(year):
    if year not in _pdf_cache:
        with pdfplumber.open(RAW / str(year) / PDF_FILES[year]) as pdf:
            _pdf_cache[year] = [p.extract_text() or "" for p in pdf.pages]
    return _pdf_cache[year]


def _find_page(texts, pattern):
    for i, t in enumerate(texts):
        if re.search(pattern, nfkc(t)):
            return i
    raise KeyError(pattern)


def _split_name_nums(line):
    """行を (名称トークン列, 先頭からの数値トークン列, 末尾側の残り) に分ける。"""
    toks = nfkc(line).split()
    name, nums, i = [], [], 0
    while i < len(toks) and not (_INT.match(toks[i])):
        name.append(toks[i]); i += 1
    while i < len(toks) and _INT.match(toks[i]):
        nums.append(toks[i]); i += 1
    return name, nums, toks[i:]


def pdf_table5(year):
    texts = pdf_texts(year)
    i0 = _find_page(texts, r"第5表\s*市町村別観光客入込数")
    left = [l for l in texts[i0].split("\n")]
    right = [l for l in texts[i0 + 1].split("\n")]
    L = []
    for l in left:
        name, nums, _ = _split_name_nums(l)
        if name and len(nums) == 6:
            L.append((" ".join(name), [to_num(x) for x in nums]))
    R = []
    for l in right:
        toks = nfkc(l).split()
        nums = []
        for t in toks:
            if _INT.match(t):
                nums.append(t)
            else:
                break
        if len(nums) >= 7:
            R.append([to_num(x) for x in nums[:7]])
    rows = []
    for (name, m1), m2 in zip(L, R):
        kind, key = norm_muni(name)
        rows.append((kind, key, {"months": m1 + m2[:6], "total": m2[6]}))
        if kind == "total":
            break
    if len(L) < 1 or not any(k == "total" for k, _, _ in rows):
        raise ValueError(f"{year} 第5表の合計行が見つかりません")
    return _group_regions(rows)


# 令和2年 第6表 南魚沼市: 「その他」(index 5) が空欄（魚沼地域計のその他=0 と整合）
SIX_BLANK_FIX = {(2020, "南魚沼市"): 5}


def pdf_table6(year):
    texts = pdf_texts(year)
    i0 = _find_page(texts, r"第6表\s*市町村別観光客入込数")
    rows = []
    for l in texts[i0].split("\n"):
        name, nums, _ = _split_name_nums(l)
        if name and len(nums) == 7 and (year, compact(" ".join(name))) in SIX_BLANK_FIX:
            # 空欄セルで7値になる行。欠けた位置（0）を補う（地域計の列合計で検算済み）
            nums = nums[:SIX_BLANK_FIX[(year, compact(" ".join(name)))]] + ["0"] + \
                nums[SIX_BLANK_FIX[(year, compact(" ".join(name)))]:]
        if name and len(nums) == 8:
            kind, key = norm_muni(" ".join(name))
            rows.append((kind, key, {"cats": [to_num(x) for x in nums[:7]], "total": to_num(nums[7])}))
            if kind == "total":
                break
    return _group_regions(rows)


def pdf_table1(year):
    texts = pdf_texts(year)
    i0 = _find_page(texts, r"第1表\s*観光客入込数の推移")
    t = nfkc(texts[i0])
    pat = re.compile(r"((?:平成|令和)\s*(?:\d+|元)年(?:/令和元年)?)\s*([0-9,]+)\s*人?\s*([0-9.]+)\s*[％%]?")
    out = {}
    for m in pat.finditer(t):
        y = era_to_year(m.group(1))
        if y:
            out[y] = (float(m.group(2).replace(",", "")), float(m.group(3)))
    if 2019 not in out:  # 「平成31年/令和元年」が改行で分断されるレイアウト（令和元年版）
        m = re.search(r"平成31年/\s*平成\d+年\s*[0-9,]+\s*[0-9.]+\s*([0-9,]+)\s*([0-9.]+)\s*令和元年", t)
        if m:
            out[2019] = (float(m.group(1).replace(",", "")), float(m.group(2)))
    return out


def table1(year):
    return excel_table1(year) if year in EXCEL_FILES else pdf_table1(year)


def table5(year):
    return excel_table5(year) if year in EXCEL_FILES else pdf_table5(year)


def table6(year):
    return excel_table6(year) if year in EXCEL_FILES else pdf_table6(year)


# ---------------- 第7表（観光地点） ----------------
def _spot_name(s):
    return re.sub(r"\s+", "", nfkc(s).replace("\n", ""))


def _cat_split(s):
    """『都市型観光 －買物・食－』→ ('都市型観光', '買物・食')。"""
    t = nfkc(s).replace("\n", " ").strip()
    m = re.match(r"(\S+)\s*[－-]\s*(.+?)\s*[－-]$", t)
    if m:
        return m.group(1), m.group(2)
    return re.sub(r"\s+", "", t), ""


def _fix_cat(spot, cat):
    """幅の狭い列で『…以外)』等の文字が分類セルに混入する場合を直す（例: '外ス)ポーツ・レクリエーション'）。"""
    if cat in CATEGORIES:
        return spot, cat
    plain = nfkc(cat).replace(")", "")
    for c in CATEGORIES:
        if plain.endswith(c):
            left = plain[:-len(c)]
            if re.fullmatch(r"-?\d+-?", left):  # ページ番号の混入
                left = ""
            return spot + left + (")" if ")" in nfkc(cat) else ""), c
    return spot, cat


def _spot_record(muni, spot, cat, cur, prev, extra=""):
    spot, cat = _fix_cat(_spot_name(spot), cat)
    cur_n, prev_n = to_num(cur), to_num(prev)
    note = []
    if prev_n is None:
        raw = compact(str(prev)) if prev is not None else ""
        note.append(f"前年:{raw}" if raw else "前年:なし")
    if extra:
        note.append(extra)
    return {"municipality": muni, "spot": _spot_name(spot), "category": cat,
            "visitors": cur_n, "prev_year_visitors": prev_n, "note": ";".join(note)}


def excel_table7(year):
    ws = _xl_sheet(year, "表7")
    hdr = None
    for r in ws.iter_rows(min_row=1, max_row=8):
        for c in r:
            if c.value and compact(str(c.value)) == "市町村":
                hdr = c
    c0 = hdr.column - 1  # 0始まり
    out, muni = [], None
    for r in ws.iter_rows(min_row=hdr.row + 1, values_only=True):
        name, spot, cat, cur, prev = r[c0], r[c0 + 1], r[c0 + 4], r[c0 + 5], r[c0 + 6]
        if name:
            muni = compact(str(name))
        if not spot or not isinstance(cur, (int, float)):
            continue
        out.append(_spot_record(muni, spot, compact(str(cat)), cur, prev))
    return out


def pdf_table7(year):
    texts = pdf_texts(year)
    i0 = _find_page(texts, r"第7表\s*市町村別主要観光地点入込数")
    out, muni = [], None
    with pdfplumber.open(RAW / str(year) / PDF_FILES[year]) as pdf:
        for pi in range(i0, len(pdf.pages)):
            pg = pdf.pages[pi]
            if not re.search(r"市町村名\s*観光地点名", nfkc(texts[pi])):
                continue
            for t in pg.find_tables():
                for row in t.extract():
                    if len(row) < 5:
                        continue
                    if row[0] and compact(row[0]) not in ("市町村名",):
                        muni = norm_muni(row[0])[1]
                    if not row[1] or "観光地点名" in compact(row[1]):
                        continue
                    cat, sub = _cat_split(row[2] or "")
                    out.append(_spot_record(muni, row[1], cat, row[3], row[4], f"副分類:{sub}" if sub else ""))
    return out


def table7(year):
    return excel_table7(year) if year in EXCEL_FILES else pdf_table7(year)

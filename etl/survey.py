"""P0: 全年の第5表(市町村別月別)/第7表(観光地点)の存在・行数を機械的に調べる。"""
import glob, json, re, sys
from pathlib import Path
import openpyxl, pdfplumber

ROOT = Path(__file__).resolve().parent.parent
NUM = re.compile(r"^[△▲＋+－-]?\d[\d,]*(\.\d+)?%?$")
MUNIS = "村上市 胎内市 関川村 粟島浦村 新潟市 新発田市 燕市 阿賀野市 聖籠町 阿賀町 弥彦村 五泉市 長岡市 三条市 柏崎市 小千谷市 加茂市 見附市 田上町 出雲崎町 刈羽村 十日町市 魚沼市 湯沢町 南魚沼市 津南町 上越市 糸魚川市 妙高市 佐渡市".split()

def norm(s): return re.sub(r"[\s　]", "", s or "")

def pdf_survey(path):
    with pdfplumber.open(path) as pdf:
        texts = [p.extract_text() or "" for p in pdf.pages]
    t5 = [i for i, t in enumerate(texts) if re.search(r"第\s*[5５7７]\s*表\s*市町村別観光客入込数（月別）", t)]
    res = {"p5_pages": [], "munis": 0, "has_total": False, "spots_rows": 0}
    if t5:
        i0 = t5[0]
        pages = texts[i0:i0 + 2]
        res["p5_pages"] = [i0 + 1, i0 + 2]
        full = norm("\n".join(pages))
        found = [m for m in MUNIS if m in full or (m == "佐渡市" and "佐渡" in full)]
        res["munis"] = len(found)
        res["missing_munis"] = [m for m in MUNIS if m not in found]
        res["has_total"] = "合計" in full
    # 第7表 行数（数値3つ以上を含む行を地点行とみなす）
    s = [i for i, t in enumerate(texts) if re.search(r"第\s*[7７9９]\s*表\s*市町村別主要観光地点入込数", t)]
    if s:
        rows = 0
        for t in texts[s[0]:]:
            for ln in t.split("\n"):
                toks = ln.split()
                nums = [x for x in toks if NUM.match(x)]
                if len(nums) >= 3 and not ln.startswith("-"):
                    rows += 1
        res["spots_rows"] = rows
    return res

def xlsx_survey(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"sheets": [(ws.title, ws.dimensions) for ws in wb]}
    ws = wb["表5市町村合計"]
    muni = [r[0].value for r in ws.iter_rows(min_row=5, max_row=45, max_col=1)
            if r[0].value and "計" not in norm(str(r[0].value)) and norm(str(r[0].value)) != ""]
    out["munis"] = len(muni)
    ws7 = wb["表7"]
    col = 6 if ws7["F5"].value == "分類" else 5   # 2022は列がA始まり
    rows = 0
    for r in ws7.iter_rows(min_row=6, values_only=True):
        vals = [v for v in r if v is not None]
        if any(isinstance(v, (int, float)) and v > 1000 for v in vals):
            rows += 1
    out["spots_rows"] = rows
    return out

def main():
    rep = {}
    for m in json.load(open(ROOT / "data/raw/manifest.json")):
        y = m["year"]
        d = ROOT / m["dir"]
        r = {}
        for x in glob.glob(str(d / "*.xlsx")):
            r["xlsx"] = xlsx_survey(x)
        for p in glob.glob(str(d / "*.pdf")):
            r["pdf"] = pdf_survey(p); r["pdf_file"] = Path(p).name
        rep[y] = r
        print(y, json.dumps(r, ensure_ascii=False), flush=True)
    (ROOT / "reports" / "p0_survey.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1))

if __name__ == "__main__":
    main()
